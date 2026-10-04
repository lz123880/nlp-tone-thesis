# -*- coding: utf-8 -*-
"""
07_export_docx.py —— Word 版论文导出模块
========================================
输入：thesis/thesis.md（论文 Markdown 源）+ output/figures/*.png
输出：thesis/thesis.docx（含封面、目录、正文、图表、参考文献）

Markdown 支持子集：# / ## / ### 标题、段落、**加粗**、*斜体*、
`行内代码`、无序/有序列表、| 表格 |、![图注](相对路径)。
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "thesis", "thesis.md")
DST = os.path.join(ROOT, "thesis", "thesis.docx")

FONT = "宋体"
FONT_HEAD = "黑体"


def set_font(run, name=FONT, size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_para(doc, text, size=12, bold=False, align=None, indent=True, name=FONT, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if indent:
        pf.first_line_indent = Pt(size * 2)
    for seg, style in parse_inline(text):
        r = p.add_run(seg)
        set_font(r, name=name, size=size, bold=bold or style == "b")
        if style == "i":
            r.italic = True
        if style == "code":
            set_font(r, name="Consolas", size=size - 1)
    return p


def parse_inline(text: str):
    """拆分 **加粗** / *斜体* / `代码`。"""
    out, buf, i = [], "", 0
    pat = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)")
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], "n"))
        if m.group(2) is not None:
            out.append((m.group(2), "b"))
        elif m.group(3) is not None:
            out.append((m.group(3), "i"))
        else:
            out.append((m.group(4), "code"))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], "n"))
    return out


def add_table(doc, header: list[str], body: list[list[str]]):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_font(r, size=10.5, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in body:
        cells = t.add_row().cells
        for j, v in enumerate(row[:len(header)]):
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(v))
            set_font(r, size=10.5)
    return t


def add_heading(doc, text, level):
    sizes = {1: 16, 2: 14, 3: 12.5}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, name=FONT_HEAD, size=sizes.get(level, 12), bold=True)


def main():
    text = open(SRC, encoding="utf-8").read()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)

    lines = text.split("\n")
    i, in_toc, cover_done = 0, False, False
    toc_items = []
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue

        # 表格
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in ln.strip("|").split("|")]
            body = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                body.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            add_table(doc, header, body)
            i = j
            continue

        # 图片
        m = re.match(r"!\[(.*?)\]\((.+?)\)", ln.strip())
        if m:
            img = os.path.normpath(os.path.join(ROOT, "thesis", m.group(2)))
            if not os.path.exists(img):
                img = os.path.normpath(os.path.join(ROOT, m.group(2)))
            if os.path.exists(img):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(img, width=Inches(5.6))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                set_font(r, size=10.5, bold=True)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            level, title = len(m.group(1)), m.group(2).strip()
            if level == 1 and not cover_done:
                # 封面
                for _ in range(5):
                    doc.add_paragraph()
                tp = doc.add_paragraph()
                tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = tp.add_run(title)
                set_font(r, name=FONT_HEAD, size=22, bold=True)
                doc.add_paragraph()
                cover_done = True
                # 收集目录项（#{1,3} 必须加括号才有捕获组）
                for l2 in lines:
                    m2 = re.match(r"^(#{1,3})\s+(.*)$", l2.strip())
                    if m2 and m2.group(2).strip() != title:
                        lv = len(m2.group(1))
                        if lv <= 2:
                            toc_items.append((lv, m2.group(2).strip()))
            else:
                add_heading(doc, title, min(level, 3))
            i += 1
            continue

        # 目录占位
        if ln.strip() == "[TOC]":
            add_heading(doc, "目  录", 1)
            for lv, t in toc_items:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(0 if lv == 1 else 24)
                r = p.add_run(t)
                set_font(r, size=12, bold=(lv == 1))
            doc.add_page_break()
            i += 1
            continue

        # 列表
        if re.match(r"^[-*]\s+", ln.strip()) or re.match(r"^\d+[.、]\s+", ln.strip()):
            add_para(doc, ln.strip(), size=12, indent=False, space_after=3)
            i += 1
            continue

        add_para(doc, ln.strip())
        i += 1

    doc.save(DST)
    print(f"Word 版已生成 → {DST}")


if __name__ == "__main__":
    main()
